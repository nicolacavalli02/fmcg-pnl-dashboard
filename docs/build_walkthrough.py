"""Build an Excel walkthrough of the FMCG P&L dataset.

The point is comprehension, not export: every figure outside the raw fact
sheet is a live formula reading the facts, so the reader can watch the P&L
ladder assemble itself instead of taking totals on trust.
"""

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(r"C:\Users\Samsung\Desktop\finance dashboards")
OUT = ROOT / "docs" / "dataset-walkthrough.xlsx"

data = json.loads((ROOT / "data" / "pnl_dataset.json").read_text("utf-8"))
meta, facts = data["meta"], data["facts"]

LINES = {l["id"]: l for l in meta["pnl_lines"]}
CUSTOMERS = {c["id"]: c for c in meta["customers"]}
CATEGORIES = {c["id"]: c for c in meta["categories"]}
CHANNELS = {c["id"]: c for c in meta["channels"]}
MONTHS = {m["n"]: m for m in meta["months"]}
CASE_LINES = {"volume_baseline", "volume_incremental", "volume_on_promo",
              "market_volume"}

# --- house style -----------------------------------------------------------
FONT = "Arial"
INK = "1F2933"
MUTED = "6B7A88"
TEAL = "0B5563"
OCHRE = "8A5D12"
BLUE = "0000FF"          # hardcoded input, by convention
RULE = Side(style="thin", color="D5DCE3")

F_TITLE = Font(name=FONT, size=15, bold=True, color=TEAL)
F_H1 = Font(name=FONT, size=11, bold=True, color="FFFFFF")
F_BODY = Font(name=FONT, size=10, color=INK)
F_MUTED = Font(name=FONT, size=9, color=MUTED, italic=True)
F_BOLD = Font(name=FONT, size=10, bold=True, color=INK)
F_STRONG = Font(name=FONT, size=10, bold=True, color=TEAL)
F_INPUT = Font(name=FONT, size=10, color=BLUE)
F_SECTION = Font(name=FONT, size=11, bold=True, color=TEAL)

FILL_HEAD = PatternFill("solid", fgColor=TEAL)
FILL_BAND = PatternFill("solid", fgColor="EEF3F5")
FILL_NOTE = PatternFill("solid", fgColor="FFF6E5")

MONEY = '#,##0;(#,##0);"-"'
MONEY2 = '#,##0.00;(#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
PP = '+0.0"pp";-0.0"pp";"-"'
CASES = '#,##0;(#,##0);"-"'
RATIO = '0.00'

wb = Workbook()


def sheet(title):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    return ws


def put(ws, cell, value, font=F_BODY, fmt=None, align=None, fill=None,
        literal=False):
    """Write a cell. `literal` keeps descriptive text that starts with "="
    from being stored as a formula, which is how openpyxl reads it by
    default."""
    c = ws[cell]
    c.value = value
    if literal and isinstance(value, str):
        c.data_type = "s"
    c.font = font
    if fmt:
        c.number_format = fmt
    if align:
        c.alignment = Alignment(horizontal=align, vertical="center")
    if fill:
        c.fill = fill
    return c


def header_row(ws, row, labels, widths=None):
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.font = F_H1
        c.fill = FILL_HEAD
        c.alignment = Alignment(horizontal="left" if i == 1 else "right",
                                vertical="center", wrap_text=True)
        c.border = Border(bottom=RULE)
    ws.row_dimensions[row].height = 28
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


# ===========================================================================
# 1. Read me
# ===========================================================================
ws = sheet("Read me")
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 26
ws.column_dimensions["C"].width = 96

put(ws, "B2", "FMCG P&L dataset - how it is put together", F_TITLE)
put(ws, "B3", "Synthetic data. The company, its customers and its results do "
    "not exist. Figures are generated from FMCG benchmark ratios.", F_MUTED)

blocks = [
    ("What is in here",
     "The current year of the dataset: Actual 2026 and Budget 2026. The other "
     "two scenarios (LY Actual 2025, Forecast 2026) are identical in shape and "
     "live in data/pnl_dataset.csv."),
    ("The grain",
     "One row per scenario, month, customer, category and P&L line. Nothing is "
     "pre-aggregated. 9,660 rows per scenario, 19,320 in this workbook."),
    ("Positive magnitudes",
     "Every value is stored positive, including costs. Whether a line adds to "
     "or subtracts from the result is carried separately, in the sign column "
     "(+1 income, -1 cost, 0 statistical). This keeps the fact table free of "
     "sign juggling and puts the arithmetic in one place."),
    ("Subtotals are not stored",
     "Net Revenue, Gross Profit, Contribution Margin, EBITDA and EBIT do not "
     "exist as rows anywhere in the data. They are derived. The 'P&L' sheet "
     "rebuilds them with live SUMIFS so you can watch it happen - that is the "
     "same job js/logic/aggregate.js does in the dashboard."),
    ("Volumes are not P&L lines",
     "Cases sold sit in the same table but outside the ladder, with sign 0. "
     "They are split three ways: baseline (what would have sold anyway), "
     "incremental (what the promotion created) and on-promo (what went out "
     "under promotional terms). Promotional pressure needs the third; "
     "promotional ROI needs the second. One measure could not answer both."),
    ("Where the customer P&L stops",
     "G&A, R&D and D&A are held per category only, with no customer - shown "
     "here as '(unallocated)'. So a customer view can reach Contribution "
     "Margin and no further. The 'Worked example' sheet shows what goes wrong "
     "if you ignore this."),
    ("Reading the colours",
     "Blue figures are raw data, typed in from the JSON. Black figures are "
     "formulas. Every total in this workbook recalculates from the Facts_CY "
     "sheet - change a value there and watch it flow."),
]

row = 5
for title, body in blocks:
    put(ws, f"B{row}", title, F_SECTION)
    c = put(ws, f"C{row}", body, F_BODY)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row].height = 15 * (len(body) // 95 + 1) + 6
    row += 1

row += 1
put(ws, f"B{row}", "Sheets", F_SECTION)
row += 1
tour = [
    ("Facts_CY", "The raw fact table. Columns A-F are exactly the JSON "
     "fields; G-L are attributes looked up from the metadata to make "
     "pivoting possible."),
    ("Worked example", "One customer, one category, one month traced from "
     "gross sales to contribution margin, then to price per case and "
     "promotional return."),
    ("P&L", "The full ladder, Actual against Budget, built with SUMIFS."),
    ("By customer", "The commercial cut, and where it has to stop."),
    ("By category", "The same cut one level up, where market share and the "
     "overhead lines become available again."),
    ("Dimensions", "Every member of every dimension, and the definition of "
     "each P&L line."),
]
for name, body in tour:
    put(ws, f"B{row}", name, F_BOLD)
    c = put(ws, f"C{row}", body, F_BODY)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row].height = 15 * (len(body) // 95 + 1) + 4
    row += 1

# ===========================================================================
# 2. Facts_CY  (referenced by every formula, so the name stays formula-safe)
# ===========================================================================
fs = sheet("Facts_CY")
rows = [f for f in facts if f["scenario"] == "cy_actual"]
rows += [f for f in facts if f["scenario"] == "budget"]

fs.merge_cells("A1:F1")
put(fs, "A1", "  Exactly the fields stored in the JSON fact table",
    Font(name=FONT, size=9, bold=True, color="FFFFFF"), fill=FILL_HEAD)
fs.merge_cells("G1:L1")
put(fs, "G1", "  Looked up from the metadata - not stored per fact",
    Font(name=FONT, size=9, bold=True, color=MUTED), fill=FILL_BAND)

header_row(fs, 2,
           ["scenario", "month", "customer", "category", "line", "value",
            "month name", "channel", "area", "line group", "sign", "unit"],
           [14, 8, 16, 15, 20, 15, 12, 15, 16, 14, 7, 8])

first = 3
for i, f in enumerate(rows):
    r = first + i
    cust = CUSTOMERS.get(f["customer"])
    line = LINES[f["line"]]
    values = [
        f["scenario"], f["month"], f["customer"] or "(unallocated)",
        f["category"], f["line"], f["value"],
        MONTHS[f["month"]]["short"],
        cust["channel"] if cust else "(unallocated)",
        cust["area"] if cust else "(unallocated)",
        line["group"], line["sign"],
        "cases" if f["line"] in CASE_LINES else "EUR",
    ]
    for col, v in enumerate(values, start=1):
        c = fs.cell(row=r, column=col, value=v)
        c.font = F_INPUT if col == 6 else F_BODY
        if col == 6:
            c.number_format = MONEY2
        if col in (2, 11):
            c.alignment = Alignment(horizontal="right")
last = first + len(rows) - 1

fs.freeze_panes = "A3"
fs.auto_filter.ref = f"A2:L{last}"

# Range constants every other sheet builds on.
SCN = f"Facts_CY!$A${first}:$A${last}"
MON = f"Facts_CY!$B${first}:$B${last}"
CUS = f"Facts_CY!$C${first}:$C${last}"
CAT = f"Facts_CY!$D${first}:$D${last}"
LIN = f"Facts_CY!$E${first}:$E${last}"
VAL = f"Facts_CY!$F${first}:$F${last}"


def sumifs(line_id, scenario, extra=""):
    return f'=SUMIFS({VAL},{SCN},"{scenario}",{LIN},"{line_id}"{extra})'


# ===========================================================================
# 3. Worked example
# ===========================================================================
ex = sheet("Worked example")
CUST, CATG, MTH = "ipervia", "beverages", 7
extra = f',{CUS},"{CUST}",{CAT},"{CATG}",{MON},{MTH}'

for col, w in zip("ABCDEF", [34, 46, 16, 16, 14, 30]):
    ex.column_dimensions[col].width = w

put(ex, "A1", "One cell, traced end to end", F_TITLE)
put(ex, "A2", f"{CUSTOMERS[CUST]['label']} - {CATEGORIES[CATG]['label']} - "
    f"{MONTHS[MTH]['label']} 2026. Every figure is a SUMIFS filtered to those "
    "three dimensions.", F_MUTED)

header_row(ex, 4, ["", "How it is computed", "Actual", "Budget", "Unit", ""],
           [34, 46, 16, 16, 14, 30])

r = 5
ref = {}


def trace(label, how, formula_a, formula_b, fmt, unit, strong=False, note=""):
    global r
    put(ex, f"A{r}", label, F_STRONG if strong else F_BODY)
    put(ex, f"B{r}", how, F_MUTED, literal=True)
    put(ex, f"C{r}", formula_a, F_STRONG if strong else F_BODY, fmt, "right")
    put(ex, f"D{r}", formula_b, F_STRONG if strong else F_BODY, fmt, "right")
    put(ex, f"E{r}", unit, F_MUTED, align="right")
    if note:
        c = put(ex, f"F{r}", note, F_MUTED, literal=True)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ex.row_dimensions[r].height = 26
    if strong:
        for col in "ABCDE":
            ex[f"{col}{r}"].border = Border(top=RULE)
    r += 1
    return r - 1


put(ex, f"A{r}", "VOLUME", F_SECTION); r += 1
ref["vb"] = trace("Baseline volume", "SUMIFS on line = volume_baseline",
                  sumifs("volume_baseline", "cy_actual", extra),
                  sumifs("volume_baseline", "budget", extra), CASES, "cases")
ref["vi"] = trace("Incremental volume", "SUMIFS on line = volume_incremental",
                  sumifs("volume_incremental", "cy_actual", extra),
                  sumifs("volume_incremental", "budget", extra), CASES, "cases",
                  note="Created by the promotion")
ref["vt"] = trace("Total volume", f"= baseline + incremental",
                  f"=C{ref['vb']}+C{ref['vi']}", f"=D{ref['vb']}+D{ref['vi']}",
                  CASES, "cases", strong=True)
ref["vp"] = trace("of which sold on promotion", "SUMIFS on line = volume_on_promo",
                  sumifs("volume_on_promo", "cy_actual", extra),
                  sumifs("volume_on_promo", "budget", extra), CASES, "cases",
                  note="Includes baseline cases that were discounted anyway")

r += 1
put(ex, f"A{r}", "GROSS TO NET", F_SECTION); r += 1
ref["gs"] = trace("Gross Sales", "SUMIFS on line = gross_sales",
                  sumifs("gross_sales", "cy_actual", extra),
                  sumifs("gross_sales", "budget", extra), MONEY, "EUR",
                  note="= cases x list price x channel index")
disc_first = r
for lid in ("discount_on_invoice", "trade_spend_promo", "trade_spend_other",
            "returns"):
    trace("   " + LINES[lid]["label"], f"SUMIFS on line = {lid}",
          sumifs(lid, "cy_actual", extra), sumifs(lid, "budget", extra),
          MONEY, "EUR")
disc_last = r - 1
ref["disc"] = trace("Discounts & Trade Spend",
                    f"= SUM of the four lines above",
                    f"=SUM(C{disc_first}:C{disc_last})",
                    f"=SUM(D{disc_first}:D{disc_last})", MONEY, "EUR")
ref["nr"] = trace("NET REVENUE", "= Gross Sales - Discounts",
                  f"=C{ref['gs']}-C{ref['disc']}",
                  f"=D{ref['gs']}-D{ref['disc']}", MONEY, "EUR", strong=True,
                  note="Not stored anywhere: derived")

r += 1
put(ex, f"A{r}", "MARGIN", F_SECTION); r += 1
cogs_first = r
for lid in ("cogs_materials", "cogs_production", "cogs_logistics"):
    trace("   " + LINES[lid]["label"], f"SUMIFS on line = {lid}",
          sumifs(lid, "cy_actual", extra), sumifs(lid, "budget", extra),
          MONEY, "EUR")
cogs_last = r - 1
ref["cogs"] = trace("Cost of Goods Sold", "= SUM of the three lines above",
                    f"=SUM(C{cogs_first}:C{cogs_last})",
                    f"=SUM(D{cogs_first}:D{cogs_last})", MONEY, "EUR",
                    note="Driven by cases x cost per case, never by revenue")
ref["gp"] = trace("GROSS PROFIT", "= Net Revenue - COGS",
                  f"=C{ref['nr']}-C{ref['cogs']}",
                  f"=D{ref['nr']}-D{ref['cogs']}", MONEY, "EUR", strong=True)
od_first = r
for lid in ("opex_marketing", "opex_sales"):
    trace("   " + LINES[lid]["label"], f"SUMIFS on line = {lid}",
          sumifs(lid, "cy_actual", extra), sumifs(lid, "budget", extra),
          MONEY, "EUR")
od_last = r - 1
ref["od"] = trace("Direct Commercial Costs", "= SUM of the two lines above",
                  f"=SUM(C{od_first}:C{od_last})",
                  f"=SUM(D{od_first}:D{od_last})", MONEY, "EUR")
ref["cm"] = trace("CONTRIBUTION MARGIN", "= Gross Profit - Direct costs",
                  f"=C{ref['gp']}-C{ref['od']}",
                  f"=D{ref['gp']}-D{ref['od']}", MONEY, "EUR", strong=True,
                  note="As far as a customer view can go")

# --- the trap --------------------------------------------------------------
r += 1
put(ex, f"A{r}", "WHERE IT HAS TO STOP", F_SECTION); r += 1
warn = r
trace("G&A, filtered to this customer", "SUMIFS on line = opex_gna",
      sumifs("opex_gna", "cy_actual", extra),
      sumifs("opex_gna", "budget", extra), MONEY, "EUR")
for cell in (f"A{warn}", f"B{warn}", f"C{warn}", f"D{warn}", f"E{warn}",
             f"F{warn}"):
    ex[cell].fill = FILL_NOTE
note = ("Excel answers 0, because no row matches - overhead is held per "
        "category with no customer. But the honest answer is 'not available "
        "at this grain', which is a different statement. Subtract that 0 and "
        "you would report EBITDA equal to Contribution Margin. This is why "
        "js/logic returns null here instead of a number, and why the ladder "
        "marks every row below it unavailable.")
ex.merge_cells(f"A{r}:F{r+2}")
c = put(ex, f"A{r}", note, Font(name=FONT, size=9, color=OCHRE), fill=FILL_NOTE)
c.alignment = Alignment(wrap_text=True, vertical="top")
r += 4

put(ex, f"A{r}", "DERIVED, PER CASE", F_SECTION); r += 1
trace("Gross price per case", "= Gross Sales / total volume",
      f"=C{ref['gs']}/C{ref['vt']}", f"=D{ref['gs']}/D{ref['vt']}",
      MONEY2, "EUR/case")
trace("Net price per case", "= Net Revenue / total volume",
      f"=C{ref['nr']}/C{ref['vt']}", f"=D{ref['nr']}/D{ref['vt']}",
      MONEY2, "EUR/case", note="The gap to gross price is the whole "
      "gross-to-net story, per case")
trace("Cost per case", "= COGS / total volume",
      f"=C{ref['cogs']}/C{ref['vt']}", f"=D{ref['cogs']}/D{ref['vt']}",
      MONEY2, "EUR/case")
gppc = trace("Gross profit per case", "= Gross Profit / total volume",
             f"=C{ref['gp']}/C{ref['vt']}", f"=D{ref['gp']}/D{ref['vt']}",
             MONEY2, "EUR/case")

r += 1
put(ex, f"A{r}", "DERIVED, COMMERCIAL", F_SECTION); r += 1
trace("Gross-to-net", "= Discounts / Gross Sales",
      f"=C{ref['disc']}/C{ref['gs']}", f"=D{ref['disc']}/D{ref['gs']}",
      PCT, "%")
trace("Gross margin", "= Gross Profit / Net Revenue",
      f"=C{ref['gp']}/C{ref['nr']}", f"=D{ref['gp']}/D{ref['nr']}", PCT, "%")
trace("Promotional pressure", "= volume on promo / total volume",
      f"=C{ref['vp']}/C{ref['vt']}", f"=D{ref['vp']}/D{ref['vt']}", PCT, "%")
promo_row = disc_first + 1
trace("Promotional ROI",
      "= incremental volume x gross profit per case / promo spend",
      f"=C{ref['vi']}*C{gppc}/C{promo_row}",
      f"=D{ref['vi']}*D{gppc}/D{promo_row}", RATIO, "x",
      note="Below 1.00 the promotion gave away more margin than it earned back")

# ===========================================================================
# 4. P&L  (whole company)
# ===========================================================================
pl = sheet("P&L")
for col, w in zip("ABCDEFGHI", [32, 40, 7, 16, 16, 15, 11, 13, 13]):
    pl.column_dimensions[col].width = w

put(pl, "A1", "Full year P&L, Actual against Budget", F_TITLE)
put(pl, "A2", "Every figure is a live formula reading the Facts_CY sheet. "
    "The indented lines are stored in the data; the bold ones are not.",
    F_MUTED)
header_row(pl, 4, ["", "How it is computed", "Sign", "Actual 2026",
                   "Budget 2026", "Variance", "Var %", "Effect",
                   "% of Net Rev"],
           [32, 40, 7, 16, 16, 15, 11, 13, 13])

r = 5
rowof = {}


def variance_cells(ws, r):
    """Variance, variance % and whether it helps or hurts. The direction is
    not the sign of the delta: a cost going up is an adverse variance even
    though the number rises, which is exactly what the sign column encodes."""
    put(ws, f"F{r}", f"=D{r}-E{r}", F_BODY, MONEY)
    put(ws, f"G{r}", f'=IF(E{r}=0,"",D{r}/E{r}-1)', F_BODY, PCT)
    put(ws, f"H{r}",
        f'=IF(F{r}=0,"",IF($C{r}*F{r}>0,"favourable","adverse"))',
        F_BODY, align="right")


def ladder_rows(target_ws, start_row, extra_filter="", show_share=True):
    """Render the ladder. Returns {structure_id: row} and the last row used."""
    r = start_row
    at = {}
    detail = []
    for srow in meta["pnl_structure"]:
        if srow["type"] == "group" and len(srow["components"]) > 1:
            first_c = r
            for lid in srow["components"]:
                put(target_ws, f"A{r}", "   " + LINES[lid]["label"], F_BODY)
                put(target_ws, f"B{r}", f"SUMIFS on line = {lid}", F_MUTED,
                    literal=True)
                put(target_ws, f"C{r}", LINES[lid]["sign"], F_MUTED,
                    align="right")
                put(target_ws, f"D{r}",
                    sumifs(lid, "cy_actual", extra_filter), F_BODY, MONEY)
                put(target_ws, f"E{r}", sumifs(lid, "budget", extra_filter),
                    F_BODY, MONEY)
                variance_cells(target_ws, r)
                detail.append(r)
                r += 1
            put(target_ws, f"A{r}", srow["label"], F_BOLD)
            put(target_ws, f"B{r}",
                f"= SUM of the {len(srow['components'])} lines above", F_MUTED,
                literal=True)
            put(target_ws, f"C{r}", LINES[srow["components"][0]]["sign"],
                F_MUTED, align="right")
            put(target_ws, f"D{r}", f"=SUM(D{first_c}:D{r-1})", F_BOLD, MONEY)
            put(target_ws, f"E{r}", f"=SUM(E{first_c}:E{r-1})", F_BOLD, MONEY)
        elif srow["type"] == "group":
            lid = srow["components"][0]
            put(target_ws, f"A{r}", srow["label"], F_BOLD)
            put(target_ws, f"B{r}", f"SUMIFS on line = {lid}", F_MUTED,
                literal=True)
            put(target_ws, f"C{r}", LINES[lid]["sign"], F_MUTED, align="right")
            put(target_ws, f"D{r}", sumifs(lid, "cy_actual", extra_filter),
                F_BOLD, MONEY)
            put(target_ws, f"E{r}", sumifs(lid, "budget", extra_filter),
                F_BOLD, MONEY)
        else:
            parts = []
            words = []
            for op, refid in srow["formula"]:
                parts.append(f"{op}D{at[refid]}")
                words.append(f"{op} {meta_label(refid)}")
            put(target_ws, f"A{r}", srow["label"].upper(), F_STRONG)
            put(target_ws, f"B{r}", "= " + " ".join(words).lstrip("+ "),
                F_MUTED, literal=True)
            put(target_ws, f"C{r}", 1, F_MUTED, align="right")
            put(target_ws, f"D{r}", "=" + "".join(parts).lstrip("+"),
                F_STRONG, MONEY)
            put(target_ws, f"E{r}",
                "=" + "".join(p.replace("D", "E") for p in parts).lstrip("+"),
                F_STRONG, MONEY)
            for col in "ABCDEFGHI":
                target_ws[f"{col}{r}"].border = Border(top=RULE)

        variance_cells(target_ws, r)
        at[srow["id"]] = r
        r += 1
    if show_share:
        nr = at["net_revenue"]
        for rr in list(at.values()) + detail:
            put(target_ws, f"I{rr}", f"=D{rr}/D${nr}", F_BODY, PCT)
    return at, r


def meta_label(refid):
    for s in meta["pnl_structure"]:
        if s["id"] == refid:
            return s["label"]
    return refid


rowof, r = ladder_rows(pl, 5)

r += 1
put(pl, f"A{r}", "DERIVED", F_SECTION); r += 1
derived = [
    ("Total volume", "= baseline + incremental volume",
     f'=SUMIFS({VAL},{SCN},"cy_actual",{LIN},"volume_baseline")'
     f'+SUMIFS({VAL},{SCN},"cy_actual",{LIN},"volume_incremental")',
     f'=SUMIFS({VAL},{SCN},"budget",{LIN},"volume_baseline")'
     f'+SUMIFS({VAL},{SCN},"budget",{LIN},"volume_incremental")', CASES),
    ("Gross-to-net", "= Discounts / Gross Sales",
     f"=D{rowof['discounts']}/D{rowof['gross_sales']}",
     f"=E{rowof['discounts']}/E{rowof['gross_sales']}", PCT),
    ("Gross margin", "= Gross Profit / Net Revenue",
     f"=D{rowof['gross_profit']}/D{rowof['net_revenue']}",
     f"=E{rowof['gross_profit']}/E{rowof['net_revenue']}", PCT),
    ("EBITDA margin", "= EBITDA / Net Revenue",
     f"=D{rowof['ebitda']}/D{rowof['net_revenue']}",
     f"=E{rowof['ebitda']}/E{rowof['net_revenue']}", PCT),
]
for label, how, fa, fb, fmt in derived:
    put(pl, f"A{r}", label, F_BODY)
    put(pl, f"B{r}", how, F_MUTED, literal=True)
    put(pl, f"D{r}", fa, F_BODY, fmt)
    put(pl, f"E{r}", fb, F_BODY, fmt)
    if fmt == PCT:
        put(pl, f"F{r}", f"=(D{r}-E{r})*100", F_BODY, PP)
    else:
        put(pl, f"F{r}", f"=D{r}-E{r}", F_BODY, CASES)
    put(pl, f"G{r}", f'=IF(E{r}=0,"",D{r}/E{r}-1)', F_BODY, PCT)
    r += 1

pl.freeze_panes = "A5"

# ===========================================================================
# 5 & 6. By customer / by category
# ===========================================================================
def cut_sheet(title, dimension, members, label_of, key_col, note,
              with_market=False):
    ws = sheet(title)
    cols = ["", "Channel" if dimension == "customer" else "Volume (cases)",
            "Gross Sales", "Discounts", "Net Revenue", "COGS", "Gross Profit",
            "Direct costs", "Contribution Margin", "Gross-to-net",
            "Gross margin", "CM %"]
    if with_market:
        cols += ["Market share, volume"]
    widths = [24, 16] + [15] * (len(cols) - 2)
    put(ws, "A1", title, F_TITLE)
    put(ws, "A2", note, F_MUTED)
    ws.merge_cells(f"A2:{get_column_letter(len(cols))}2")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 30
    header_row(ws, 4, cols, widths)

    r = 5
    for m in members:
        ex_f = f',{key_col},"{m}"'
        put(ws, f"A{r}", label_of(m), F_BODY)
        if dimension == "customer":
            put(ws, f"B{r}", CHANNELS[CUSTOMERS[m]["channel"]]["label"],
                F_MUTED, align="right")
        else:
            put(ws, f"B{r}",
                f'=SUMIFS({VAL},{SCN},"cy_actual",{LIN},"volume_baseline"{ex_f})'
                f'+SUMIFS({VAL},{SCN},"cy_actual",{LIN},"volume_incremental"{ex_f})',
                F_BODY, CASES)
        put(ws, f"C{r}", sumifs("gross_sales", "cy_actual", ex_f), F_BODY, MONEY)
        put(ws, f"D{r}", "=" + "+".join(
            sumifs(l, "cy_actual", ex_f)[1:] for l in
            ("discount_on_invoice", "trade_spend_promo", "trade_spend_other",
             "returns")), F_BODY, MONEY)
        put(ws, f"E{r}", f"=C{r}-D{r}", F_BOLD, MONEY)
        put(ws, f"F{r}", "=" + "+".join(
            sumifs(l, "cy_actual", ex_f)[1:] for l in
            ("cogs_materials", "cogs_production", "cogs_logistics")),
            F_BODY, MONEY)
        put(ws, f"G{r}", f"=E{r}-F{r}", F_BOLD, MONEY)
        put(ws, f"H{r}", "=" + "+".join(
            sumifs(l, "cy_actual", ex_f)[1:] for l in
            ("opex_marketing", "opex_sales")), F_BODY, MONEY)
        put(ws, f"I{r}", f"=G{r}-H{r}", F_BOLD, MONEY)
        put(ws, f"J{r}", f"=D{r}/C{r}", F_BODY, PCT)
        put(ws, f"K{r}", f"=G{r}/E{r}", F_BODY, PCT)
        put(ws, f"L{r}", f"=I{r}/E{r}", F_BODY, PCT)
        if with_market:
            put(ws, f"M{r}",
                f'=(SUMIFS({VAL},{SCN},"cy_actual",{LIN},"volume_baseline"{ex_f})'
                f'+SUMIFS({VAL},{SCN},"cy_actual",{LIN},"volume_incremental"{ex_f}))'
                f'/SUMIFS({VAL},{SCN},"cy_actual",{LIN},"market_volume"{ex_f})',
                F_BODY, PCT)
        r += 1

    put(ws, f"A{r}", "Total", F_STRONG)
    for col in "CDEFGHI":
        put(ws, f"{col}{r}", f"=SUM({col}5:{col}{r-1})", F_STRONG, MONEY)
        ws[f"{col}{r}"].border = Border(top=RULE)
    for col, num in (("J", f"=D{r}/C{r}"), ("K", f"=G{r}/E{r}"),
                     ("L", f"=I{r}/E{r}")):
        put(ws, f"{col}{r}", num, F_STRONG, PCT)
        ws[f"{col}{r}"].border = Border(top=RULE)
    ws[f"A{r}"].border = Border(top=RULE)
    ws[f"B{r}"].border = Border(top=RULE)
    ws.freeze_panes = "B5"
    return ws, r


cut_sheet("By customer", "customer", list(CUSTOMERS), lambda m: CUSTOMERS[m]["label"],
          CUS,
          "The commercial cut. It stops at Contribution Margin on purpose: "
          "overhead is not held per customer, so EBITDA cannot be stated here "
          "without inventing an allocation key. Note how Discount contributes "
          "less than Wholesale despite discounting less - the list price is "
          "lower to begin with.")

cut_sheet("By category", "category", list(CATEGORIES), lambda m: CATEGORIES[m]["label"],
          CAT,
          "The same cut one level up. Here the market lines exist, so share "
          "becomes computable - and Dairy is where the volume shortfall turns "
          "into a share loss rather than a soft market.",
          with_market=True)

# ===========================================================================
# 7. Dimensions
# ===========================================================================
dm = sheet("Dimensions")
for col, w in zip("ABCDEF", [28, 20, 20, 18, 10, 18]):
    dm.column_dimensions[col].width = w
put(dm, "A1", "Dimensions and line definitions", F_TITLE)

r = 3
put(dm, f"A{r}", "Customers", F_SECTION); r += 1
header_row(dm, r, ["Customer", "Channel", "Area", "", "", ""],
           [28, 20, 20, 18, 10, 18])
r += 1
for c in meta["customers"]:
    put(dm, f"A{r}", c["label"], F_BODY)
    put(dm, f"B{r}", CHANNELS[c["channel"]]["label"], F_BODY, align="right")
    put(dm, f"C{r}", c["area"], F_BODY, align="right")
    r += 1

r += 1
put(dm, f"A{r}", "Categories", F_SECTION); r += 1
header_row(dm, r, ["Category", "", "", "", "", ""])
r += 1
for c in meta["categories"]:
    put(dm, f"A{r}", c["label"], F_BODY)
    r += 1

r += 1
put(dm, f"A{r}", "P&L lines", F_SECTION); r += 1
header_row(dm, r, ["Line id", "Label", "Group", "Sign", "Per customer?", ""],
           [28, 34, 20, 18, 14, 18])
dm.column_dimensions["B"].width = 34
r += 1
for l in meta["pnl_lines"]:
    put(dm, f"A{r}", l["id"], F_BODY)
    put(dm, f"B{r}", l["label"], F_BODY, align="right")
    put(dm, f"C{r}", l["group"], F_MUTED, align="right")
    put(dm, f"D{r}", l["sign"], F_BODY, align="right")
    put(dm, f"E{r}", "yes" if l["customer_level"] else "no - category only",
        F_BODY if l["customer_level"] else Font(name=FONT, size=10,
                                                color=OCHRE), align="right")
    r += 1

del wb["Sheet"]
wb.active = 0
OUT.parent.mkdir(exist_ok=True)
wb.save(OUT)
print(f"written {OUT}  ({OUT.stat().st_size / 1024:.0f} KB)")
print(f"fact rows {first}..{last}  ({len(rows):,})")
